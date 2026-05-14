#!/usr/bin/env python3
"""Ingest Excel rows and selected web pages into a Milvus knowledge collection."""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import math
import os
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
import warnings
from collections import deque
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import Iterable

warnings.filterwarnings("ignore", message="pkg_resources is deprecated as an API.*", category=UserWarning)

try:
    import openpyxl
except ImportError:  # pragma: no cover - dependency guard
    openpyxl = None


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_EXCEL = (
    "/Users/cyc/Library/Containers/com.tencent.xinWeChat/Data/Documents/xwechat_files/"
    "wxid_t3me8j3lw09h22_9ac8/temp/RWTemp/2026-04/"
    "37f597022a6dd4e3fa2b44e5833b0c6f/氢脆应力应变曲线抽取要求-A0.xlsx"
)
DEFAULT_DB_PATH = str(ROOT / "data" / "milvus_knowledge.db")
DEFAULT_COLLECTION = "szlab_knowledge"
DEFAULT_SITE_URL = "https://std.samr.gov.cn/"
DEFAULT_ALLOWED_DOMAINS = ["std.samr.gov.cn", "openstd.samr.gov.cn"]
DEFAULT_MARKDOWN = (
    "/Users/cyc/Library/Containers/com.tencent.xinWeChat/Data/Documents/xwechat_files/"
    "wxid_t3me8j3lw09h22_9ac8/msg/file/2026-05/材料大词典第二版.md"
)
METADATA_SCHEMA_VERSION = "2026-05-14.1"


@dataclass
class Document:
    text: str
    title: str
    source_type: str
    source_uri: str
    metadata: dict


class TextAndLinkParser(HTMLParser):
    def __init__(self, base_url: str):
        super().__init__()
        self.base_url = base_url
        self.title = ""
        self.links: list[str] = []
        self._skip_depth = 0
        self._in_title = False
        self._chunks: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attrs_map = dict(attrs)
        if tag in {"script", "style", "noscript", "svg"}:
            self._skip_depth += 1
            return
        if tag == "title":
            self._in_title = True
        if tag == "a" and attrs_map.get("href"):
            self.links.append(urllib.parse.urljoin(self.base_url, attrs_map["href"]))
        if tag in {"p", "br", "div", "section", "article", "li", "tr", "h1", "h2", "h3", "h4"}:
            self._chunks.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript", "svg"} and self._skip_depth:
            self._skip_depth -= 1
        if tag == "title":
            self._in_title = False
        if tag in {"p", "div", "section", "article", "li", "tr", "h1", "h2", "h3", "h4"}:
            self._chunks.append("\n")

    def handle_data(self, data: str) -> None:
        if self._skip_depth:
            return
        value = html.unescape(data).strip()
        if not value:
            return
        if self._in_title:
            self.title = value
        self._chunks.append(value)

    def get_text(self) -> str:
        text = " ".join(self._chunks)
        text = re.sub(r"[ \t\r\f\v]+", " ", text)
        text = re.sub(r"\n\s+", "\n", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()


def load_env_file(path: Path) -> None:
    if not path.exists():
        return
    for line in path.read_text(encoding="utf-8").splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith("#") or "=" not in stripped:
            continue
        key, raw_value = stripped.split("=", 1)
        os.environ.setdefault(key.strip(), raw_value.strip().strip("\"'"))


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\u0000", " ").strip()
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    return text


def split_metadata_list(value: object) -> list[str]:
    text = normalize_text(value)
    if not text:
        return []
    return [
        item.strip()
        for item in re.split(r"\s*(?:[|/、,，;；]|\bor\b|\band\b)\s*", text, flags=re.I)
        if item.strip()
    ]


def compact_metadata(metadata: dict) -> dict:
    compacted = {}
    for key, value in metadata.items():
        if value in (None, "", [], {}):
            continue
        compacted[key] = value
    return compacted


def source_id(*parts: object) -> str:
    return ":".join(normalize_text(part).replace(":", "_") for part in parts if normalize_text(part))


def base_metadata(
    *,
    source_type: str,
    source_id_value: str,
    canonical_name: str = "",
    aliases: list[str] | None = None,
    definition: str = "",
    field_name: str = "",
    unit: str = "",
    data_type: str = "",
    section: str = "",
    evidence_type: str = "",
    source_title: str = "",
    source_path: str = "",
    source_url: str = "",
    extra: dict | None = None,
) -> dict:
    return compact_metadata(
        {
            "schema_version": METADATA_SCHEMA_VERSION,
            "source_id": source_id_value,
            "source_type": source_type,
            "source_title": source_title,
            "source_path": source_path,
            "source_url": source_url,
            "canonical_name": canonical_name,
            "aliases": [item for item in aliases or [] if item and item != canonical_name],
            "definition": definition,
            "field_name": field_name,
            "unit": unit,
            "data_type": data_type,
            "section": section,
            "evidence_type": evidence_type,
            **(extra or {}),
        }
    )


def find_header_row(rows: list[list[object]]) -> int:
    for index, row in enumerate(rows):
        normalized = [normalize_text(cell) for cell in row]
        if "列名" in normalized and "定义" in normalized:
            return index
    return 0


def row_to_document(
    workbook_path: Path,
    sheet_name: str,
    row_number: int,
    headers: list[str],
    values: list[object],
    current_section: str,
    current_field: str,
) -> Document | None:
    pairs = []
    row_fields = {}

    for index, raw_value in enumerate(values):
        value = normalize_text(raw_value)
        if not value:
            continue
        header = headers[index] if index < len(headers) and headers[index] else f"列{index + 1}"
        row_fields[header] = value
        pairs.append(f"{header}: {value}")

    if not pairs:
        return None

    def field_value(*names: str) -> str:
        for name in names:
            if name in row_fields:
                return row_fields[name]
        for header, value in row_fields.items():
            if any(name in header for name in names):
                return value
        return ""

    definition = field_value("定义")
    aliases = split_metadata_list(field_value("涵盖参数", "同义词", "别名", "英文别名"))
    data_type = field_value("数据类型", "类型")
    unit = field_value("单位")
    canonical_name = current_field or field_value("列名", "字段")
    evidence_type = "excel_includes_parameter" if aliases else "excel_row"
    metadata = base_metadata(
        source_type="excel",
        source_id_value=source_id("excel", workbook_path.name, sheet_name, row_number),
        canonical_name=canonical_name,
        aliases=aliases,
        definition=definition,
        field_name=canonical_name,
        unit=unit,
        data_type=data_type,
        section=current_section,
        evidence_type=evidence_type,
        source_title=workbook_path.name,
        source_path=str(workbook_path),
        extra={
            "workbook": workbook_path.name,
            "sheet": sheet_name,
            "row_number": row_number,
            "raw_fields": row_fields,
        },
    )

    text = "\n".join(
        [
            f"来源: {workbook_path.name}",
            f"工作表: {sheet_name}",
            f"原始行号: {row_number}",
            f"章节/类别: {current_section or '未标注'}",
            f"字段: {current_field or '延续上一字段'}",
            *pairs,
        ]
    )
    title = f"{workbook_path.name} / {sheet_name} / row {row_number}"
    return Document(
        text=text,
        title=title,
        source_type="excel",
        source_uri=f"{workbook_path}#{sheet_name}!{row_number}",
        metadata=metadata,
    )


def load_excel_documents(paths: Iterable[str]) -> list[Document]:
    if openpyxl is None:
        raise RuntimeError("openpyxl is not installed. Run: pip install -r requirements-rag.txt")

    documents: list[Document] = []
    for raw_path in paths:
        workbook_path = Path(raw_path).expanduser()
        if not workbook_path.exists():
            raise FileNotFoundError(f"Excel file not found: {workbook_path}")
        workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
        for worksheet in workbook.worksheets:
            rows = list(worksheet.iter_rows(values_only=True))
            if not rows:
                continue
            header_index = find_header_row(rows[:20])
            headers = [normalize_text(cell) or f"列{idx + 1}" for idx, cell in enumerate(rows[header_index])]
            current_section = ""
            current_field = ""
            for offset, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
                values = list(row)
                first_cell = normalize_text(values[0] if values else "")
                second_cell = normalize_text(values[1] if len(values) > 1 else "")
                if first_cell:
                    current_section = first_cell
                if second_cell:
                    current_field = second_cell
                document = row_to_document(
                    workbook_path=workbook_path,
                    sheet_name=worksheet.title,
                    row_number=offset,
                    headers=headers,
                    values=values,
                    current_section=current_section,
                    current_field=current_field,
                )
                if document:
                    documents.append(document)
    return documents


def clean_markdown_block(text: str) -> str:
    cleaned = re.sub(r"!\[[^\]]*]\([^)]+\)", "", text)
    cleaned = re.sub(r"<t[dh][^>]*>(.*?)</t[dh]>", lambda match: f"\n{html.unescape(match.group(1)).strip()}\n", cleaned, flags=re.I | re.S)
    cleaned = re.sub(r"<[^>]+>", " ", cleaned)
    cleaned = html.unescape(cleaned)
    cleaned = re.sub(r"[ \t\r\f\v]+", " ", cleaned)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip()


def load_markdown_documents(paths: Iterable[str], start_heading: str = "", stop_pattern: str = "") -> list[Document]:
    documents: list[Document] = []
    heading_re = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
    start_heading = normalize_text(start_heading)
    stop_re = re.compile(stop_pattern) if stop_pattern else None

    for raw_path in paths:
        markdown_path = Path(raw_path).expanduser()
        if not markdown_path.exists():
            raise FileNotFoundError(f"Markdown file not found: {markdown_path}")

        all_lines = markdown_path.read_text(encoding="utf-8", errors="replace").splitlines()
        if start_heading:
            start_index = next(
                (
                    index
                    for index, line in enumerate(all_lines)
                    if (match := heading_re.match(line)) and start_heading in normalize_text(match.group(2))
                ),
                None,
            )
            if start_index is None:
                raise ValueError(f"Markdown start heading not found in {markdown_path}: {start_heading}")
            line_offset = start_index
            lines = all_lines[start_index:]
        else:
            line_offset = 0
            lines = all_lines
        if stop_re:
            stop_index = next((index for index, line in enumerate(lines) if stop_re.search(line)), None)
            if stop_index is not None:
                lines = lines[:stop_index]

        heading_stack: list[str] = []
        buffer: list[str] = []
        section_start = line_offset + 1

        def emit(end_line: int) -> None:
            nonlocal buffer, section_start
            content = clean_markdown_block("\n".join(buffer))
            if len(content) < 40:
                buffer = []
                return
            section = " / ".join(heading_stack) or markdown_path.stem
            text = "\n".join(
                [
                    f"来源: {markdown_path.name}",
                    f"章节: {section}",
                    f"原始行号: {section_start}-{end_line}",
                    content,
                ]
            )
            documents.append(
                Document(
                    text=text,
                    title=f"{markdown_path.name} / {section}",
                    source_type="markdown",
                    source_uri=f"{markdown_path}#L{section_start}",
                    metadata=base_metadata(
                        source_type="markdown",
                        source_id_value=source_id("markdown", markdown_path.name, section_start),
                        canonical_name=section,
                        section=section,
                        evidence_type="markdown_section",
                        source_title=markdown_path.name,
                        source_path=str(markdown_path),
                        extra={
                            "file": markdown_path.name,
                            "start_line": section_start,
                            "end_line": end_line,
                        },
                    ),
                )
            )
            buffer = []

        for local_line_number, line in enumerate(lines, start=1):
            line_number = line_offset + local_line_number
            heading_match = heading_re.match(line)
            if heading_match:
                emit(line_number - 1)
                level = len(heading_match.group(1))
                heading = normalize_text(heading_match.group(2))
                heading_stack = heading_stack[: level - 1]
                heading_stack.append(heading)
                section_start = line_number
                buffer = [line]
                continue

            if not buffer:
                section_start = line_number
            buffer.append(line)

        emit(len(lines))

    return documents


def parse_markdown_entry_start(line: str) -> dict | None:
    line = clean_markdown_block(line)
    line = re.sub(r"^#{1,6}\s+", "", line).strip()
    line = re.sub(r"^\|?[-:| ]+\|?$", "", line).strip()
    if not line or line.startswith("[") or len(line) < 4:
        return None
    if re.fullmatch(r"[A-Z0-9Ｏ0其他\s.．-]{1,16}", line):
        return None

    reference_match = re.match(r"^(.{1,36}?)(参见|见)(.{1,90})[。.]?$", line)
    if reference_match and re.search(r"[\u4e00-\u9fff]", reference_match.group(1)):
        term = reference_match.group(1).strip()
        target = reference_match.group(3).strip()
        if "。" not in term and "，" not in term and len(target) >= 2:
            return {
                "term": term,
                "english": "",
                "definition": f"{reference_match.group(2)}{target}",
                "kind": "reference",
            }

    for match in re.finditer(r"[A-Za-z]", line):
        start = match.start()
        term_part = line[:start].strip()
        if not re.search(r"[\u4e00-\u9fff]", term_part) or len(term_part) > 60:
            continue
        definition_match = re.search(r"[\u4e00-\u9fff]", line[start:])
        if not definition_match:
            continue
        definition_start = start + definition_match.start()
        english = line[start:definition_start].strip(" ;；:：")
        definition = line[definition_start:].strip()
        if len(english) < 2 or len(definition) < 2:
            continue

        last_cjk = max(index for index, char in enumerate(term_part) if "\u4e00" <= char <= "\u9fff")
        english_prefix = term_part[last_cjk + 1 :].strip()
        term = term_part[: last_cjk + 1].strip()
        if english_prefix and not re.search(r"[\u4e00-\u9fff]", english_prefix):
            english = f"{english_prefix}{english}"
        if not term or len(term) > 48:
            continue
        return {"term": term, "english": english, "definition": definition, "kind": "definition"}

    return None


def markdown_term_aliases(term: str, english: str, definition: str, kind: str) -> tuple[list[str], str]:
    aliases = []
    if english:
        aliases.append(english)
    evidence_type = "dictionary_see_also" if kind == "reference" else "dictionary_definition"

    alias_match = re.match(r"^又称(.+?)[。；;，,]", definition)
    if alias_match:
        aliases.extend(split_metadata_list(alias_match.group(1)))
        evidence_type = "dictionary_alias"

    see_also_match = re.match(r"^(?:见|参见)(.+?)(?:[（(]|[。；;，,]|$)", definition)
    if see_also_match:
        aliases.append(see_also_match.group(1).strip())
        evidence_type = "dictionary_see_also"

    aliases = [item for item in dict.fromkeys(aliases) if item and item != term]
    return aliases, evidence_type


def load_markdown_term_documents(paths: Iterable[str], start_heading: str = "", stop_pattern: str = "") -> list[Document]:
    documents: list[Document] = []
    heading_re = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
    start_heading = normalize_text(start_heading)
    stop_re = re.compile(stop_pattern) if stop_pattern else None

    for raw_path in paths:
        markdown_path = Path(raw_path).expanduser()
        if not markdown_path.exists():
            raise FileNotFoundError(f"Markdown file not found: {markdown_path}")

        all_lines = markdown_path.read_text(encoding="utf-8", errors="replace").splitlines()
        if start_heading:
            start_index = next(
                (
                    index
                    for index, line in enumerate(all_lines)
                    if (match := heading_re.match(line)) and start_heading in normalize_text(match.group(2))
                ),
                None,
            )
            if start_index is None:
                raise ValueError(f"Markdown start heading not found in {markdown_path}: {start_heading}")
        else:
            start_index = 0

        body_index = next(
            (
                index
                for index in range(start_index, len(all_lines))
                if re.match(r"^#{1,6}\s*A\s*$", all_lines[index].strip())
            ),
            start_index,
        )
        end_index = len(all_lines)
        for index, line in enumerate(all_lines[body_index:], start=body_index):
            heading_match = heading_re.match(line)
            heading = normalize_text(heading_match.group(2)) if heading_match else ""
            if heading == "英文索引" or (stop_re and stop_re.search(line)):
                end_index = index
                break

        current: dict | None = None
        current_lines: list[str] = []
        current_start_line = 0
        current_section = ""

        def emit(end_line: int) -> None:
            nonlocal current, current_lines, current_start_line
            if not current:
                return
            continuation = clean_markdown_block("\n".join(current_lines)).strip()
            definition = current["definition"]
            if continuation:
                definition = f"{definition}\n{continuation}"
            definition = clean_markdown_block(definition)
            if len(definition) < 3:
                current = None
                current_lines = []
                return

            aliases, evidence_type = markdown_term_aliases(current["term"], current["english"], definition, current["kind"])
            english_line = f"英文: {current['english']}" if current["english"] else "英文: 未标注"
            text = "\n".join(
                [
                    f"来源: {markdown_path.name}",
                    f"术语: {current['term']}",
                    english_line,
                    f"章节: {current_section or '正文'}",
                    f"原始行号: {current_start_line}-{end_line}",
                    f"定义: {definition}",
                ]
            )
            title_suffix = f" / {current['english']}" if current["english"] else ""
            documents.append(
                Document(
                    text=text,
                    title=f"{current['term']}{title_suffix}",
                    source_type="markdown_term",
                    source_uri=f"{markdown_path}#L{current_start_line}",
                    metadata=base_metadata(
                        source_type="markdown_term",
                        source_id_value=source_id("markdown_term", markdown_path.name, current_start_line),
                        canonical_name=current["term"],
                        aliases=aliases,
                        definition=definition,
                        section=current_section,
                        evidence_type=evidence_type,
                        source_title=markdown_path.name,
                        source_path=str(markdown_path),
                        extra={
                            "file": markdown_path.name,
                            "term": current["term"],
                            "english": current["english"],
                            "kind": current["kind"],
                            "start_line": current_start_line,
                            "end_line": end_line,
                        },
                    ),
                )
            )
            current = None
            current_lines = []

        for index in range(body_index, end_index):
            line_number = index + 1
            raw_line = all_lines[index]
            stripped = raw_line.strip()
            if not stripped or stripped.startswith("![]("):
                continue

            heading_match = heading_re.match(stripped)
            if heading_match:
                heading = normalize_text(heading_match.group(2))
                parsed_heading = parse_markdown_entry_start(stripped)
                if parsed_heading:
                    emit(line_number - 1)
                    current = parsed_heading
                    current_lines = []
                    current_start_line = line_number
                    continue
                if re.fullmatch(r"[A-Z0-9Ｏ0其他\s.．-]{1,16}", heading):
                    emit(line_number - 1)
                    current_section = heading
                    continue

            parsed = parse_markdown_entry_start(raw_line)
            if parsed:
                emit(line_number - 1)
                current = parsed
                current_lines = []
                current_start_line = line_number
                continue

            if current:
                current_lines.append(raw_line)

        emit(end_index)

    return documents


def is_allowed_url(url: str, allowed_domains: set[str]) -> bool:
    parsed = urllib.parse.urlparse(url)
    if parsed.scheme not in {"http", "https"}:
        return False
    host = parsed.netloc.lower()
    return host in allowed_domains


def clean_url(url: str) -> str:
    parsed = urllib.parse.urlparse(url)
    path = parsed.path
    if path == "/":
        path = ""
    elif path.endswith("/"):
        path = path.rstrip("/")
    parsed = parsed._replace(netloc=parsed.netloc.lower(), path=path, fragment="")
    return urllib.parse.urlunparse(parsed)


def fetch_html(url: str, timeout: int) -> tuple[str, str]:
    request = urllib.request.Request(
        url,
        headers={
            "User-Agent": "szlab-rag-ingestor/0.1 (+local knowledge base ingestion)",
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        },
    )
    with urllib.request.urlopen(request, timeout=timeout) as response:
        content_type = response.headers.get("Content-Type", "")
        if "html" not in content_type.lower():
            return "", content_type
        charset = response.headers.get_content_charset() or "utf-8"
        body = response.read().decode(charset, errors="replace")
        return body, content_type


def crawl_site_documents(
    start_url: str,
    allowed_domains: set[str],
    max_pages: int,
    max_depth: int,
    timeout: int,
    delay: float,
) -> tuple[list[Document], dict]:
    queue: deque[tuple[str, int]] = deque([(clean_url(start_url), 0)])
    visited: set[str] = set()
    documents: list[Document] = []
    stats = {"visited": 0, "fetched": 0, "failed": 0, "skipped_non_html": 0}

    while queue and len(visited) < max_pages:
        url, depth = queue.popleft()
        url = clean_url(url)
        if url in visited or not is_allowed_url(url, allowed_domains):
            continue
        visited.add(url)
        stats["visited"] += 1

        try:
            body, _content_type = fetch_html(url, timeout=timeout)
        except (urllib.error.URLError, TimeoutError, UnicodeError):
            stats["failed"] += 1
            continue
        if not body:
            stats["skipped_non_html"] += 1
            continue
        stats["fetched"] += 1

        parser = TextAndLinkParser(url)
        parser.feed(body)
        text = parser.get_text()
        if len(text) >= 80:
            documents.append(
                Document(
                    text=text,
                    title=parser.title or url,
                    source_type="web",
                    source_uri=url,
                    metadata=base_metadata(
                        source_type="web",
                        source_id_value=source_id("web", url),
                        canonical_name=parser.title or url,
                        definition=text[:500],
                        evidence_type="web_page",
                        source_title=parser.title or url,
                        source_url=url,
                        extra={"url": url, "depth": depth, "title": parser.title or ""},
                    ),
                )
            )

        if depth < max_depth:
            for link in parser.links:
                cleaned = clean_url(link)
                if cleaned not in visited and is_allowed_url(cleaned, allowed_domains):
                    queue.append((cleaned, depth + 1))

        if delay:
            time.sleep(delay)

    return documents, stats


def split_text(text: str, max_chars: int, overlap: int) -> list[str]:
    paragraphs = [item.strip() for item in re.split(r"\n{2,}|(?<=[。！？.!?])\s+", text) if item.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        if len(paragraph) > max_chars:
            if current:
                chunks.append(current.strip())
                current = ""
            for start in range(0, len(paragraph), max_chars - overlap):
                chunks.append(paragraph[start : start + max_chars].strip())
            continue
        candidate = f"{current}\n{paragraph}".strip() if current else paragraph
        if len(candidate) <= max_chars:
            current = candidate
        else:
            chunks.append(current.strip())
            current = paragraph
    if current:
        chunks.append(current.strip())
    return [chunk for chunk in chunks if chunk]


def tokenize(text: str) -> list[str]:
    lowered = text.lower()
    words = re.findall(r"[a-z0-9][a-z0-9_\-/.]{1,}", lowered)
    cjk = re.findall(r"[\u4e00-\u9fff]", text)
    cjk_grams = ["".join(cjk[index : index + 2]) for index in range(max(0, len(cjk) - 1))]
    return words + cjk + cjk_grams


def hash_embedding(text: str, dimension: int) -> list[float]:
    vector = [0.0] * dimension
    for token in tokenize(text):
        digest = hashlib.blake2b(token.encode("utf-8"), digest_size=8).digest()
        value = int.from_bytes(digest, "big")
        index = value % dimension
        sign = 1.0 if (value >> 63) == 0 else -1.0
        vector[index] += sign
    norm = math.sqrt(sum(item * item for item in vector)) or 1.0
    return [item / norm for item in vector]


def call_embedding_api(texts: list[str], dimension: int) -> list[list[float]]:
    base_url = os.environ.get("EMBEDDING_BASE_URL") or os.environ.get("LLM_BASE_URL", "https://api.openai.com/v1")
    api_key = os.environ.get("EMBEDDING_API_KEY") or os.environ.get("LLM_API_KEY", "")
    model = os.environ.get("EMBEDDING_MODEL", "text-embedding-3-small")
    path = os.environ.get("EMBEDDING_PATH", "/embeddings")
    if not api_key:
        raise RuntimeError("EMBEDDING_API_KEY or LLM_API_KEY is required for openai embedding mode.")
    payload = json.dumps({"model": model, "input": texts}).encode("utf-8")
    request = urllib.request.Request(
        f"{base_url.rstrip('/')}{path}",
        data=payload,
        headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
        method="POST",
    )
    with urllib.request.urlopen(request, timeout=120) as response:
        data = json.loads(response.read().decode("utf-8"))
    vectors = [item["embedding"] for item in sorted(data["data"], key=lambda item: item["index"])]
    for vector in vectors:
        if len(vector) != dimension:
            raise RuntimeError(f"Embedding dimension mismatch: expected {dimension}, got {len(vector)}")
    return vectors


def embed_texts(texts: list[str], mode: str, dimension: int, batch_size: int) -> list[list[float]]:
    vectors: list[list[float]] = []
    for start in range(0, len(texts), batch_size):
        batch = texts[start : start + batch_size]
        if mode == "openai":
            vectors.extend(call_embedding_api(batch, dimension))
        else:
            vectors.extend(hash_embedding(text, dimension) for text in batch)
    return vectors


def stable_int_id(value: str) -> int:
    digest = hashlib.blake2b(value.encode("utf-8"), digest_size=8).digest()
    return int.from_bytes(digest, "big") & ((1 << 63) - 1)


def build_records(documents: list[Document], max_chars: int, overlap: int) -> list[dict]:
    records: list[dict] = []
    for document in documents:
        for chunk_index, chunk in enumerate(split_text(document.text, max_chars=max_chars, overlap=overlap)):
            source_key = f"{document.source_uri}#{chunk_index}"
            metadata = {**document.metadata, "chunk_index": chunk_index}
            records.append(
                {
                    "id": stable_int_id(source_key),
                    "text": chunk,
                    "title": document.title[:512],
                    "source_type": document.source_type,
                    "source_uri": document.source_uri[:1024],
                    "chunk_index": chunk_index,
                    "metadata_json": json.dumps(metadata, ensure_ascii=False),
                }
            )
    return records


def connect_milvus(db_path: str):
    try:
        from pymilvus import MilvusClient
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise RuntimeError("pymilvus is not installed. Run: pip install -r requirements-rag.txt") from exc
    Path(db_path).expanduser().parent.mkdir(parents=True, exist_ok=True)
    return MilvusClient(str(Path(db_path).expanduser()))


def insert_records(client, collection_name: str, records: list[dict], dimension: int, reset: bool, batch_size: int) -> int:
    if reset and client.has_collection(collection_name=collection_name):
        client.drop_collection(collection_name=collection_name)
    if not client.has_collection(collection_name=collection_name):
        client.create_collection(collection_name=collection_name, dimension=dimension)

    inserted = 0
    for start in range(0, len(records), batch_size):
        batch = records[start : start + batch_size]
        result = client.insert(collection_name=collection_name, data=batch)
        inserted += int(result.get("insert_count", len(batch))) if isinstance(result, dict) else len(batch)
    return inserted


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Ingest Excel and web content into a Milvus Lite collection.")
    parser.add_argument("--excel", action="append", default=[], help="Excel file path. Can be passed multiple times.")
    parser.add_argument("--markdown", action="append", default=[], help="Markdown file path. Can be passed multiple times.")
    parser.add_argument("--default-markdown", action="store_true", help="Ingest the default material dictionary Markdown file.")
    parser.add_argument("--markdown-mode", choices=["section", "term"], default="section", help="Markdown ingestion mode.")
    parser.add_argument("--markdown-start-heading", default="", help="Skip Markdown content before the first heading containing this text.")
    parser.add_argument("--markdown-stop-pattern", default=r"^\[General Information\]", help="Skip Markdown content from the first line matching this regex.")
    parser.add_argument("--skip-default-excel", action="store_true", help="Do not ingest the default hydrogen embrittlement Excel file.")
    parser.add_argument("--site-url", default=DEFAULT_SITE_URL, help="Start URL for website ingestion.")
    parser.add_argument("--skip-site", action="store_true", help="Skip website crawling.")
    parser.add_argument("--allow-domain", action="append", default=[], help="Allowed crawl domain. Can be passed multiple times.")
    parser.add_argument("--site-max-pages", type=int, default=30, help="Maximum pages to crawl.")
    parser.add_argument("--site-depth", type=int, default=1, help="Maximum crawl depth from start URL.")
    parser.add_argument("--site-timeout", type=int, default=20, help="HTTP timeout in seconds.")
    parser.add_argument("--site-delay", type=float, default=0.2, help="Delay between page requests.")
    parser.add_argument("--db-path", default=os.environ.get("MILVUS_DB_PATH", DEFAULT_DB_PATH))
    parser.add_argument("--collection", default=os.environ.get("MILVUS_COLLECTION", DEFAULT_COLLECTION))
    parser.add_argument("--embedding-mode", choices=["hash", "openai"], default=os.environ.get("EMBEDDING_MODE", "hash"))
    parser.add_argument("--embedding-dim", type=int, default=int(os.environ.get("EMBEDDING_DIM", "384")))
    parser.add_argument("--chunk-chars", type=int, default=1000)
    parser.add_argument("--chunk-overlap", type=int, default=120)
    parser.add_argument("--batch-size", type=int, default=128)
    parser.add_argument("--reset", action="store_true", help="Drop and recreate the collection before inserting.")
    parser.add_argument("--dry-run", action="store_true", help="Extract and chunk sources without writing to Milvus.")
    return parser.parse_args(argv)


def main(argv: list[str]) -> int:
    load_env_file(ROOT / ".env")
    args = parse_args(argv)

    excel_paths = list(args.excel)
    if not args.skip_default_excel:
        excel_paths.insert(0, DEFAULT_EXCEL)

    documents: list[Document] = []
    source_counts = {"excel": 0, "markdown": 0, "markdown_term": 0, "web": 0}
    if excel_paths:
        excel_documents = load_excel_documents(excel_paths)
        source_counts["excel"] = len(excel_documents)
        documents.extend(excel_documents)

    markdown_paths = list(args.markdown)
    if args.default_markdown:
        markdown_paths.insert(0, DEFAULT_MARKDOWN)
    if markdown_paths:
        markdown_loader = load_markdown_term_documents if args.markdown_mode == "term" else load_markdown_documents
        markdown_documents = markdown_loader(markdown_paths, start_heading=args.markdown_start_heading, stop_pattern=args.markdown_stop_pattern)
        source_counts["markdown" if args.markdown_mode == "section" else "markdown_term"] = len(markdown_documents)
        documents.extend(markdown_documents)

    site_stats = {}
    if not args.skip_site:
        allowed_domains = set(args.allow_domain or DEFAULT_ALLOWED_DOMAINS)
        web_documents, site_stats = crawl_site_documents(
            start_url=args.site_url,
            allowed_domains=allowed_domains,
            max_pages=args.site_max_pages,
            max_depth=args.site_depth,
            timeout=args.site_timeout,
            delay=args.site_delay,
        )
        source_counts["web"] = len(web_documents)
        documents.extend(web_documents)

    records = build_records(documents, max_chars=args.chunk_chars, overlap=args.chunk_overlap)
    print(
        json.dumps(
            {
                "documents": len(documents),
                "source_counts": source_counts,
                "site_stats": site_stats,
                "records": len(records),
                "collection": args.collection,
                "db_path": args.db_path,
                "embedding_mode": args.embedding_mode,
                "embedding_dim": args.embedding_dim,
                "dry_run": args.dry_run,
            },
            ensure_ascii=False,
        )
    )

    if args.dry_run:
        for record in records[:5]:
            preview = {key: record[key] for key in ["id", "title", "source_type", "source_uri", "chunk_index"]}
            preview["text_preview"] = record["text"][:180]
            print(json.dumps(preview, ensure_ascii=False))
        return 0

    texts = [record["text"] for record in records]
    vectors = embed_texts(texts, mode=args.embedding_mode, dimension=args.embedding_dim, batch_size=args.batch_size)
    for record, vector in zip(records, vectors):
        record["vector"] = vector

    client = connect_milvus(args.db_path)
    inserted = insert_records(
        client=client,
        collection_name=args.collection,
        records=records,
        dimension=args.embedding_dim,
        reset=args.reset,
        batch_size=args.batch_size,
    )
    print(json.dumps({"inserted": inserted, "collection": args.collection, "db_path": args.db_path}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
